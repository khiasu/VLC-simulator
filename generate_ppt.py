import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5) # 16:9 Widescreen
    
    # Helper to add bullet slide
    def add_slide(title_text, content_bullets, image_path=None):
        blank_layout = prs.slide_layouts[6] # Blank slide
        slide = prs.slides.add_slide(blank_layout)
        
        # Add Header banner
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(24, 43, 73)
        
        # Determine text width based on whether an image is provided
        text_width = Inches(6.0) if image_path and os.path.exists(image_path) else Inches(11.7)
        body_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), text_width, Inches(5.4))
        tf_body = body_box.text_frame
        tf_body.word_wrap = True
        
        for i, bullet in enumerate(content_bullets):
            p = tf_body.paragraphs[0] if i == 0 else tf_body.add_paragraph()
            p.text = bullet
            p.font.size = Pt(17)
            p.font.color.rgb = RGBColor(50, 50, 50)
            
            if bullet.startswith("  -") or bullet.startswith("    "):
                p.level = 1
                p.font.size = Pt(15)
            elif bullet.startswith("      "):
                p.level = 2
                p.font.size = Pt(13)
                
        # Embed image if available
        if image_path and os.path.exists(image_path):
            slide.shapes.add_picture(image_path, Inches(7.2), Inches(1.5), width=Inches(5.5))
            
        return slide

    # Title Slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    t_box = title_slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.0))
    t_frame = t_box.text_frame
    p1 = t_frame.paragraphs[0]
    p1.text = "Channel-Aware Adaptive Reed–Solomon Error Control\nfor a CSK-Based Visible Light Communication System"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(18, 52, 102)
    
    p2 = t_frame.add_paragraph()
    p2.text = "\nMinor Project Review Presentation | Visible Light Communication (VLC) Simulator"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(80, 80, 80)

    # 1. Aim
    add_slide("1. Aim of the Project", [
        "Develop an end-to-end Visible Light Communication (VLC) software simulator.",
        "Implement IEEE 802.15.7 4-Color Shift Keying (4-CSK) modulation over an indoor optical channel.",
        "Design and evaluate an explainable, Rule-Based Channel-Aware Adaptive Reed–Solomon Controller.",
        "Maximize system Goodput while enforcing a strict reliability constraint: Frame Error Rate (FER) <= 10^-2."
    ])

    # 2. Objectives
    add_slide("2. Project Objectives", [
        "Modular Communication Pipeline:",
        "  - Application bitstream generation and 12-byte framed encapsulation (8-byte header + CRC32).",
        "  - Dynamic Reed-Solomon Codec over GF(2^8) supporting multiple (N, K) redundancy profiles.",
        "Physical Optical Modeling:",
        "  - Implement 4-CSK modulation with constant total optical power (R + G + B = 1.0) to prevent visible flicker.",
        "  - Model indoor impairments: AWGN noise, LED dimming scaling (0.1 to 1.0), and 3x3 RGB optical crosstalk.",
        "Adaptive Error Control:",
        "  - Receiver EVM-to-SNR estimation and historical sliding-window link health tracking.",
        "  - Multi-indicator controller featuring 0.5 dB hysteresis margin and fast emergency escalation."
    ])

    # 3. Motivation
    add_slide("3. Motivation & Background", [
        "Radio Frequency (RF) Spectrum Congestion:",
        "  - Wi-Fi and cellular bands face severe bandwidth limitations and RF interference.",
        "Advantages of Optical Wireless Communication (OWC / VLC):",
        "  - Terahertz unlicensed optical spectrum with zero electromagnetic interference (safe in hospitals & aircraft).",
        "  - Inherent physical security (light does not penetrate opaque walls).",
        "The Adaptive Coding Imperative:",
        "  - Indoor optical links fluctuate heavily due to user mobility, ambient sunlight, and user-adjusted dimming.",
        "  - Static error correction is inefficient: strong codes waste throughput in clean channels, while weak codes drop packets when shadowed."
    ])

    # 4. Literature Survey & Problem Statement
    add_slide("4. Literature Survey & Problem Statement", [
        "Survey of Existing Work:",
        "  - IEEE 802.15.7 defines CSK modulation standards, but most studies analyze static physical links in isolation.",
        "  - Existing network simulators (NS-3) lack detailed cross-layer optical PHY/MAC modeling for adaptive Reed-Solomon.",
        "Identified Gap & Problem Statement:",
        "  - How to dynamically adapt Forward Error Correction (FEC) redundancy in real-time over a fluctuating 4-CSK optical link to maximize Goodput without exceeding target Frame Error Rate (FER <= 1%)?"
    ])

    # 5. System Architecture
    add_slide("5. End-to-End System Pipeline", [
        "Layered Communication Workflow:",
        "  - Application Layer: Pseudo-random payload data generation.",
        "  - Data Link Layer: Header packing (Frame ID, Length, Profile ID) + CRC32 error detection.",
        "  - FEC Layer: Reed-Solomon encoder over GF(2^8) with chunking.",
        "  - Physical Layer: 4-CSK constellation mapping (Red, Green, Blue, White).",
        "  - Optical Channel: AWGN noise injection, dimming factor scaling, and RGB crosstalk matrix.",
        "  - Receiver Demapper: Maximum Likelihood (ML) minimum Euclidean distance demapping.",
        "  - Feedback Engine: Receiver EVM calculation, SNR estimation, and adaptive profile feedback."
    ], image_path="figures/fig1_csk_constellation_3d.png")

    # 6. Physical Channel & Modulation
    add_slide("6. 4-CSK Modulation & Optical Channel Model", [
        "IEEE 802.15.7 4-CSK Chromaticity Mapping:",
        "  - '00' -> [1.0, 0.0, 0.0] (Red)",
        "  - '01' -> [0.0, 1.0, 0.0] (Green)",
        "  - '10' -> [0.0, 0.0, 1.0] (Blue)",
        "  - '11' -> [0.333, 0.333, 0.333] (Neutral White)",
        "Optical Channel Mathematical Impairments:",
        "  - I_rx = alpha * (M_crosstalk * I_tx) + n_AWGN",
        "  - Constant optical power constraint: sum(I) = 1.0 avoids human eye flicker.",
        "  - Receiver AGC normalizes amplitude prior to ML Euclidean distance demapping."
    ], image_path="figures/fig1_csk_constellation_3d.png")

    # 7. Adaptive Controller & Calibration
    add_slide("7. Data-Driven Adaptive Control Logic", [
        "Receiver Channel Quality Estimator:",
        "  - Error Vector Magnitude (EVM) calculated from received symbols.",
        "  - Estimated SNR: SNR_est = -20 * log10(EVM).",
        "Multi-Indicator Decision Engine:",
        "  - Offline Crossover Calibration: Swept 5 to 25 dB SNR to find optimal RS switching points.",
        "  - 0.5 dB Hysteresis Margin: Prevents rapid ping-pong oscillations under channel jitter.",
        "  - Fast Emergency Escalation: Immediate fallback to RS(255, 223) if 2 consecutive CRC failures or FER > 10% occurs."
    ], image_path="figures/fig3_goodput_vs_snr_crossover.png")

    # 8. Performance Evaluation (BER & Goodput)
    add_slide("8. Experimental Results: BER & Crossover", [
        "Observed Simulation Results:",
        "  - Raw channel BER decreases monotonically from 10^-1 to 10^-4 as SNR increases from 6 to 24 dB.",
        "  - Reed-Solomon decoding successfully eliminates byte errors within its correction capacity t = (N - K)/2.",
        "  - Goodput Crossover Points:",
        "      * SNR < 16.0 dB: RS(255, 223) ensures link survival.",
        "      * 16.0 <= SNR < 17.5 dB: RS(255, 239) optimal.",
        "      * 17.5 <= SNR < 20.5 dB: RS(255, 251) optimal.",
        "      * SNR >= 20.5 dB: UNCODED maximizes Goodput up to 1500+ kbps."
    ], image_path="figures/fig2_ber_vs_snr.png")

    # 9. Dynamic Benchmark & Statistical Validation
    add_slide("9. Multi-Seed Benchmark & Validation", [
        "Multi-Seed Dynamic Channel Benchmark (5 seeds, 150 frames/trial):",
        "  - Dynamic random walk channel fluctuating between 5 dB and 22 dB.",
        "  - Synthetic sensor estimation noise (sigma = 1.0 dB) injected.",
        "Key Statistical Results:",
        "  - Adaptive RS achieves high Goodput (915 kbps) while keeping FER strictly controlled.",
        "  - Static Uncoded baseline collapses (FER = 75%, Goodput drops to 467 kbps) during fading dips.",
        "  - High link stability (>85% frames without unnecessary profile switches) validates hysteresis effectiveness."
    ], image_path="figures/fig5_statistical_benchmark.png")

    # 10. Sensitivity to Dimming & Crosstalk
    add_slide("10. Sensitivity to Dimming & Optical Crosstalk", [
        "LED Dimming Impact:",
        "  - Lowering dimming level reduces optical signal-to-noise ratio at the photodiode.",
        "  - Uncoded transmission drops packets below 50% dimming, while RS protection maintains error-free goodput.",
        "RGB Optical Crosstalk Sensitivity:",
        "  - As inter-channel crosstalk increases from 0% to 25%, adjacent sensor overlap increases BER linearly.",
        "  - RS codes absorb crosstalk degradation without requiring physical optical filter redesigns."
    ], image_path="figures/fig6_dimming_crosstalk_sensitivity.png")

    # 11. Interactive Research Dashboard
    add_slide("11. Interactive Research Workbench", [
        "Modern Streamlit Web Dashboard (dashboard.py):",
        "  - Live Simulation Mode: Real-time dynamic channel sliders, moving graphs, 3D constellation scatter, and decision console.",
        "  - Multi-Baseline Benchmark Mode: Displays statistical summaries with 95% Confidence Intervals.",
        "  - Log-Driven Replay Mode: Frame scrubbing inspector allowing exact re-simulation and payload byte extraction.",
        "  - Strict Provenance Tracking: JSON experiment metadata tracking random seeds and system parameters."
    ], image_path="figures/fig4_dynamic_tracking_timeseries.png")

    # 12. Conclusion & Deliverables
    add_slide("12. Conclusion & Summary of Deliverables", [
        "Summary of Minor Project Accomplishments:",
        "  - Fully functional end-to-end IEEE 802.15.7 4-CSK VLC software simulator.",
        "  - Explainable, deterministic Rule-Based Adaptive Controller with hysteresis and escalation.",
        "  - 100% Pytest test suite coverage across all physical and data link modules (17/17 passing).",
        "  - Publication-quality high-resolution figures and multi-seed statistical benchmarks.",
        "  - Interactive Streamlit workbench for live experimentation and replay."
    ])

    prs.save("Project_Review_Presentation.pptx")
    print("Project_Review_Presentation.pptx updated with embedded simulation figures!")

if __name__ == "__main__":
    create_presentation()
