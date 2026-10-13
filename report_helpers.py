import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Sets cell padding in dxa (60 dxa = 3pt, 100 dxa = 5pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F1F5F9"):
    """Sets cell background fill."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_academic_table_borders(table):
    """
    Applies exact academic booktabs table borders matching Likhibi:
    Top rule: single, 1pt (sz=8), black (000000)
    Bottom rule: single, 1pt (sz=8), black (000000)
    Inside horizontal: single, 0.5pt (sz=4), light gray (E2E8F0)
    Vertical/left/right: None
    """
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_borderless_table(table):
    """Removes all borders for TOC, LOT, LOF, LOS, and equation tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
    f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_body_p(doc, text):
    """
    Adds standard academic body paragraph matching Likhibi:
    Times New Roman, 12 pt, 1.5 line spacing (360), Space After 6 pt (120), Justified.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pPr = p._p.get_or_add_pPr()
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="120" w:line="360" w:lineRule="auto"/>')
    pPr.append(spacing)
    jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="both"/>')
    pPr.append(jc)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    """Adds an indented bullet point with 1.5 line spacing and Times New Roman 12 pt."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pPr = p._p.get_or_add_pPr()
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="80" w:line="360" w:lineRule="auto"/>')
    pPr.append(spacing)
    ind = parse_xml(f'<w:ind {nsdecls("w")} w:left="360" w:hanging="180"/>')
    pPr.append(ind)
    
    # Bullet symbol
    r_sym = p.add_run("▪  ")
    r_sym.font.name = 'Times New Roman'
    r_sym.font.size = Pt(10)
    r_sym.font.bold = True
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        
    r_txt = p.add_run(text)
    r_txt.font.name = 'Times New Roman'
    r_txt.font.size = Pt(12)
    return p

def add_chapter_title(doc, number, title):
    """
    Adds chapter title matching Likhibi:
    CHAPTER X
    <TITLE>
    Left-aligned, Bold, Space Before 12 pt (240), Space After 12 pt (240), 1.5 lines (360).
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._p.get_or_add_pPr()
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="240" w:line="360" w:lineRule="auto"/>')
    pPr.append(spacing)
    jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="left"/>')
    pPr.append(jc)
    
    r1 = p.add_run(f"CHAPTER {number}\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(14)
    r1.font.bold = True
    
    r2 = p.add_run(title.upper())
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(14)
    r2.font.bold = True
    return p

def add_heading_1(doc, text):
    """
    Adds Heading 1 matching Likhibi (e.g. 1.1 Motivation):
    Bold, Left-aligned, Space Before 12 pt (240), Space After 4 pt (80), 1.5 lines.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._p.get_or_add_pPr()
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="80" w:line="360" w:lineRule="auto"/>')
    pPr.append(spacing)
    jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="left"/>')
    pPr.append(jc)
    
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    return p

def add_heading_2(doc, text):
    """
    Adds Heading 2 matching Likhibi (e.g. 1.1.1 Sub-topic):
    Bold, Left-aligned, Space Before 8 pt (160), Space After 3 pt (60), 1.5 lines.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p._p.get_or_add_pPr()
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="160" w:after="60" w:line="360" w:lineRule="auto"/>')
    pPr.append(spacing)
    jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="left"/>')
    pPr.append(jc)
    
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    return p

def add_table_with_caption(doc, table_num_str, caption_text, headers, rows_data, col_widths=None):
    """
    Adds academic data table with top caption matching Likhibi:
    - Caption ABOVE: Times New Roman, 11 pt Bold, Left-aligned, Space Before 8 pt (160), Space After 3 pt (60)
    - Table: Centered, booktabs horizontal borders, header fill #F1F5F9 (10.5 pt bold), data cells (10 pt regular)
    """
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p_cap._p.get_or_add_pPr()
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="160" w:after="60"/>')
    pPr.append(spacing)
    r_cap = p_cap.add_run(f"Table {table_num_str}  {caption_text}")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.bold = True

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_academic_table_borders(table)

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr_h = p._p.get_or_add_pPr()
        sp_h = parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="60"/>')
        pPr_h.append(sp_h)
        r = p.add_run(h_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.font.bold = True
        set_cell_shading(hdr_cells[i], "F1F5F9")
        set_cell_margins(hdr_cells[i], top=60, bottom=60, left=100, right=100)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if col_widths and i < len(col_widths):
            hdr_cells[i].width = Inches(col_widths[i])

    # Format Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = ""
            p = row_cells[c_idx].paragraphs[0]
            # Left align first column, center others if short/numbers
            if c_idx == 0 and len(str(val)) > 10:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pPr_d = p._p.get_or_add_pPr()
            sp_d = parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="60"/>')
            pPr_d.append(sp_d)
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if col_widths and c_idx < len(col_widths):
                row_cells[c_idx].width = Inches(col_widths[c_idx])

    # Add small spacing after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
    return table

def add_figure_with_caption(doc, img_path, caption_text, fig_num_str, width_inches=5.8):
    """
    Adds publication-grade figure with centered image and bottom caption matching Likhibi:
    - Image: Centered
    - Caption BELOW: Times New Roman, 10.5 pt Bold, Centered, Space Before 2 pt (40), Space After 10 pt (200)
    """
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p_cap._p.get_or_add_pPr()
        spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="40" w:after="200"/>')
        pPr.append(spacing)
        jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>')
        pPr.append(jc)
        r_cap = p_cap.add_run(f"Figure {fig_num_str}  {caption_text}")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10.5)
        r_cap.font.bold = True
    else:
        add_body_p(doc, f"[Figure {fig_num_str}: {caption_text} - Image file not found at {img_path}]")

def make_equation_box(doc, eq_text, eq_num_str):
    """Typesets an equation with right-aligned equation number in borderless 2-col table matching Likhibi."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borderless_table(table)
    
    # Left cell: Equation
    c_eq = table.rows[0].cells[0]
    c_eq.width = Inches(5.4)
    p_eq = c_eq.paragraphs[0]
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr_eq = p_eq._p.get_or_add_pPr()
    sp_eq = parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="60"/>')
    pPr_eq.append(sp_eq)
    r_eq = p_eq.add_run(eq_text)
    r_eq.font.name = 'Times New Roman'
    r_eq.font.size = Pt(11)
    r_eq.font.italic = True
    
    # Right cell: Equation number
    c_num = table.rows[0].cells[1]
    c_num.width = Inches(0.9)
    p_num = c_num.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr_num = p_num._p.get_or_add_pPr()
    sp_num = parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="60"/>')
    pPr_num.append(sp_num)
    r_num = p_num.add_run(f"... ({eq_num_str})")
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(11)
    r_num.font.bold = True
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def setup_footer_page_number(section, is_roman=False):
    """
    Configures right-aligned page number in footer matching Likhibi:
    - is_roman=True: 'fmt=lowerRoman', 'start=1'
    - is_roman=False: 'start=1' (decimal)
    """
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    
    # Configure pgNumType
    sectPr = section._sectPr
    # Remove existing pgNumType if any
    for existing in sectPr.xpath('./w:pgNumType'):
        sectPr.remove(existing)
        
    pgNumType = OxmlElement('w:pgNumType')
    if is_roman:
        pgNumType.set(qn('w:fmt'), 'lowerRoman')
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)
    
    # Setup footer paragraph
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.text = ""
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_pPr = f_p._p.get_or_add_pPr()
    jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="right"/>')
    f_pPr.append(jc)
    
    # Add PAGE field XML exactly like Likhibi word/footer2.xml
    fld = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'  <w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/></w:rPr>'
        f'  <w:fldChar w:fldCharType="begin"/>'
        f'</w:r>'
    )
    fld_instr = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'  <w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/></w:rPr>'
        f'  <w:instrText xml:space="preserve"> PAGE </w:instrText>'
        f'</w:r>'
    )
    fld_sep = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>')
    fld_val = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'  <w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/></w:rPr>'
        f'  <w:t>1</w:t>'
        f'</w:r>'
    )
    fld_end = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')
    
    f_p._p.append(fld)
    f_p._p.append(fld_instr)
    f_p._p.append(fld_sep)
    f_p._p.append(fld_val)
    f_p._p.append(fld_end)
