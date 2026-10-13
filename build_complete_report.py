import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from build_report_frontmatter import create_report
from report_chapters_1_3 import add_chapters_1_to_3
from report_chapters_4_6 import add_chapters_4_to_6
from report_chapters_7_8 import add_chapters_7_and_8

def add_footer_page_number(run):
    """Inserts a dynamic Word PAGE field into a run."""
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def main():
    print("Building VLC Mini Project Report for St. Joseph University...")
    
    # 1. Initialize Document & Generate Front Matter
    doc, add_p, add_heading_1, add_heading_2, add_chapter_title, add_figure_with_caption = create_report()
    
    # 2. Add Chapters 1 to 3
    print("Adding Chapters 1 to 3...")
    add_chapters_1_to_3(doc, add_p, add_heading_1, add_heading_2, add_chapter_title, add_figure_with_caption)
    
    # 3. Add Chapters 4 to 6
    print("Adding Chapters 4 to 6...")
    add_chapters_4_to_6(doc, add_p, add_heading_1, add_heading_2, add_chapter_title, add_figure_with_caption)
    
    # 4. Add Chapters 7 and 8, References, Annexures
    print("Adding Chapters 7 and 8, References, and Annexures...")
    add_chapters_7_and_8(doc, add_p, add_heading_1, add_heading_2, add_chapter_title, add_figure_with_caption)
    
    # 5. Add Running Headers & Footers
    print("Configuring running headers and footers...")
    for idx, section in enumerate(doc.sections):
        # Footer with centered dynamic page number
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ft.paragraph_format.space_before = Pt(0)
        p_ft.paragraph_format.space_after = Pt(0)
        r_ft = p_ft.add_run()
        r_ft.font.name = 'Times New Roman'
        r_ft.font.size = Pt(10)
        add_footer_page_number(r_ft)
        
        # Subtle header
        header = section.header
        p_hd = header.paragraphs[0]
        p_hd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hd.paragraph_format.space_before = Pt(0)
        p_hd.paragraph_format.space_after = Pt(0)
        r_hd = p_hd.add_run("VLC Simulator | Mini Project Report | Dept. of CSE, St. Joseph University")
        r_hd.font.name = 'Times New Roman'
        r_hd.font.size = Pt(8.5)
        r_hd.font.italic = True
        r_hd.font.color.rgb = RGBColor(120, 120, 120)

    # 6. Save the Final Report Document
    output_filename = "VLC_MINI_PROJECT_REPORT_SJU.docx"
    doc.save(output_filename)
    print(f"\n=======================================================")
    print(f"SUCCESS: Report saved as '{output_filename}'")
    print(f"File size: {os.path.getsize(output_filename):,} bytes")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    print(f"=======================================================")

if __name__ == "__main__":
    main()
