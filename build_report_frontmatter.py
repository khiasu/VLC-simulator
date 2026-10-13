import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from report_helpers import (
    set_borderless_table,
    set_cell_margins,
    set_cell_shading,
    setup_footer_page_number,
    add_body_p
)

def build_frontmatter(doc):
    """
    Builds Section 0 (Cover page + Bonafide Certificate) and Section 1 (Frontmatter)
    with exact formatting matching LIKHIBI REPORT MINI PROJECT (1).docx.
    """
    # -------------------------------------------------------------
    # SECTION 0: COVER PAGE & BONAFIDE CERTIFICATE
    # -------------------------------------------------------------
    sec0 = doc.sections[0]
    sec0.page_width = Inches(8.27)
    sec0.page_height = Inches(11.69)
    sec0.top_margin = Inches(1.0)
    sec0.bottom_margin = Inches(1.0)
    sec0.left_margin = Inches(1.0)
    sec0.right_margin = Inches(1.0)
    sec0.header_distance = Inches(0.5)
    sec0.footer_distance = Inches(0.5)

    # Paragraph 0: Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_title._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="160" w:line="320" w:lineRule="auto"/>'))
    r_title = p_title.add_run("CHANNEL-AWARE ADAPTIVE REED-SOLOMON ERROR CONTROL FOR A CSK-BASED VISIBLE LIGHT COMMUNICATION SYSTEM\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(17)
    r_title.font.bold = True

    # Paragraph 1: A MINI PROJECT REPORT
    p_rep = doc.add_paragraph()
    p_rep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_rep._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="80" w:after="160"/>'))
    r_rep = p_rep.add_run("A MINI PROJECT REPORT\n")
    r_rep.font.name = 'Times New Roman'
    r_rep.font.size = Pt(14)
    r_rep.font.bold = True

    # Paragraph 2: SUBMITTED BY
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_sub._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="100"/>'))
    r_sub = p_sub.add_run("SUBMITTED BY\n")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True

    # Paragraph 3: Candidate Name & Reg. No.
    p_cand = doc.add_paragraph()
    p_cand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_cand._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="140" w:line="300" w:lineRule="auto"/>'))
    r_c1 = p_cand.add_run("KHIASUTHONG T\n")
    r_c1.font.name = 'Times New Roman'
    r_c1.font.size = Pt(15)
    r_c1.font.bold = True
    r_c2 = p_cand.add_run("(Reg. No.: 2306107010)\n")
    r_c2.font.name = 'Times New Roman'
    r_c2.font.size = Pt(14)
    r_c2.font.bold = True

    # Paragraph 4: In Partial Fulfillment
    p_ful = doc.add_paragraph()
    p_ful.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_ful._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="100"/>'))
    r_ful = p_ful.add_run("IN PARTIAL FULFILLMENT FOR THE AWARD OF THE DEGREE OF\n")
    r_ful.font.name = 'Times New Roman'
    r_ful.font.size = Pt(13)
    r_ful.font.bold = True

    # Paragraph 5: Degree & University Logo
    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_deg._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="120"/>'))
    r_deg = p_deg.add_run("B. TECH COMPUTER SCIENCE AND ENGINEERING\n")
    r_deg.font.name = 'Times New Roman'
    r_deg.font.size = Pt(15)
    r_deg.font.bold = True

    # Insert Crest / Logo
    logo_path = os.path.join(os.path.dirname(__file__), "extracted_template_media", "image1.png")
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(8)
        p_logo.paragraph_format.space_after = Pt(12)
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Inches(2.1))

    # Paragraph 6-8: Department, University, Location
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_dept._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="40" w:line="280" w:lineRule="auto"/>'))
    r_d1 = p_dept.add_run("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
    r_d1.font.name = 'Times New Roman'
    r_d1.font.size = Pt(15)
    r_d1.font.bold = True

    r_d2 = p_dept.add_run("ST. JOSEPH UNIVERSITY\n")
    r_d2.font.name = 'Times New Roman'
    r_d2.font.size = Pt(15)
    r_d2.font.bold = True

    r_d3 = p_dept.add_run("NAGALAND - 797115\n")
    r_d3.font.name = 'Times New Roman'
    r_d3.font.size = Pt(14)
    r_d3.font.bold = True

    # Month & Year
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_date._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="120" w:after="0"/>'))
    r_date = p_date.add_run("NOVEMBER 2026")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(14)

    # Page Break to Certificate
    doc.add_page_break()

    # -------------------------------------------------------------
    # BONAFIDE CERTIFICATE PAGE
    # -------------------------------------------------------------
    p_cert_title = doc.add_paragraph()
    p_cert_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_cert_title._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="200" w:after="300"/>'))
    r_ct = p_cert_title.add_run("BONAFIDE CERTIFICATE")
    r_ct.font.name = 'Times New Roman'
    r_ct.font.size = Pt(16)
    r_ct.font.bold = True

    p_cert_body = doc.add_paragraph()
    p_cert_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pPr = p_cert_body._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="500" w:line="420" w:lineRule="auto"/>'))
    
    r_cb1 = p_cert_body.add_run('This Mini Project Report entitled "')
    r_cb1.font.name = 'Times New Roman'
    r_cb1.font.size = Pt(13)

    r_cb2 = p_cert_body.add_run('CHANNEL-AWARE ADAPTIVE REED-SOLOMON ERROR CONTROL FOR A CSK-BASED VISIBLE LIGHT COMMUNICATION SYSTEM')
    r_cb2.font.name = 'Times New Roman'
    r_cb2.font.size = Pt(13)
    r_cb2.font.bold = True

    r_cb3 = p_cert_body.add_run('" is the bonafide work of ')
    r_cb3.font.name = 'Times New Roman'
    r_cb3.font.size = Pt(13)

    r_cb4 = p_cert_body.add_run('KHIASUTHONG T (Reg. No.: 2306107010) of DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING')
    r_cb4.font.name = 'Times New Roman'
    r_cb4.font.size = Pt(13)
    r_cb4.font.bold = True

    r_cb5 = p_cert_body.add_run(' who carried out the project work under my supervision. It is certified further that to the best of our knowledge the work reported herein is not the part of any other project report or submitted for the award of any other degree or diploma.')
    r_cb5.font.name = 'Times New Roman'
    r_cb5.font.size = Pt(13)

    # Date of Viva-Voce
    p_viva = doc.add_paragraph()
    p_viva.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = p_viva._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="200" w:after="400"/>'))
    r_vv = p_viva.add_run("Date of Viva-Voce: ______________________")
    r_vv.font.name = 'Times New Roman'
    r_vv.font.size = Pt(12)
    r_vv.font.bold = True

    # Supervisor & HOD Signatures
    tbl_sigs = doc.add_table(rows=1, cols=2)
    tbl_sigs.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borderless_table(tbl_sigs)
    tbl_sigs.autofit = False

    c_sup = tbl_sigs.rows[0].cells[0]
    c_sup.width = Inches(3.3)
    p_s = c_sup.paragraphs[0]
    p_s.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_s.paragraph_format.line_spacing = 1.3
    r_s = p_s.add_run("___________________________\nSUPERVISOR\nMs. Alovi N Shohe\nAssistant Professor\nDepartment of Computer Science & Engineering\nSt. Joseph University")
    r_s.font.name = 'Times New Roman'
    r_s.font.size = Pt(11)
    r_s.font.bold = True

    c_hod = tbl_sigs.rows[0].cells[1]
    c_hod.width = Inches(3.3)
    p_h = c_hod.paragraphs[0]
    p_h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_h.paragraph_format.line_spacing = 1.3
    r_h = p_h.add_run("___________________________\nHEAD OF DEPARTMENT\nMs. Nempi\nAssistant Professor & Head\nDepartment of Computer Science & Engineering\nSt. Joseph University")
    r_h.font.name = 'Times New Roman'
    r_h.font.size = Pt(11)
    r_h.font.bold = True

    # Spacing paragraph
    p_mid = doc.add_paragraph()
    p_mid.paragraph_format.space_before = Pt(36)
    p_mid.paragraph_format.space_after = Pt(24)

    # Examiners Signatures
    tbl_ex = doc.add_table(rows=1, cols=2)
    tbl_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borderless_table(tbl_ex)
    tbl_ex.autofit = False

    c_int = tbl_ex.rows[0].cells[0]
    c_int.width = Inches(3.3)
    p_int = c_int.paragraphs[0]
    p_int.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_int = p_int.add_run("___________________________\nINTERNAL EXAMINER")
    r_int.font.name = 'Times New Roman'
    r_int.font.size = Pt(11)
    r_int.font.bold = True

    c_ext = tbl_ex.rows[0].cells[1]
    c_ext.width = Inches(3.3)
    p_ext = c_ext.paragraphs[0]
    p_ext.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ext = p_ext.add_run("___________________________\nEXTERNAL EXAMINER")
    r_ext.font.name = 'Times New Roman'
    r_ext.font.size = Pt(11)
    r_ext.font.bold = True

    # -------------------------------------------------------------
    # SECTION 1: FRONTMATTER (Roman Numerals: i, ii, iii...)
    # -------------------------------------------------------------
    sec1 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    sec1.page_width = Inches(8.27)
    sec1.page_height = Inches(11.69)
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.0)
    sec1.right_margin = Inches(1.0)
    sec1.header_distance = Inches(0.5)
    sec1.footer_distance = Inches(0.5)
    setup_footer_page_number(sec1, is_roman=True)

    # ACKNOWLEDGEMENT
    p_ack = doc.add_paragraph()
    p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_ack._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="360"/>'))
    r_at = p_ack.add_run("ACKNOWLEDGEMENT")
    r_at.font.name = 'Times New Roman'
    r_at.font.size = Pt(14)
    r_at.font.bold = True

    add_body_p(doc, "First and foremost, I offer my profound gratitude to Almighty God for His divine grace, wisdom, health, and sustaining providence throughout the conceptualization, system modeling, simulation execution, and successful completion of this mini project.")

    add_body_p(doc, "I express my deepest sense of gratitude and sincere indebtedness to my respected project supervisor, Ms. Alovi N Shohe, Assistant Professor, Department of Computer Science & Engineering, St. Joseph University, Nagaland, for her patient mentorship, incisive technical suggestions, rigorous evaluations, and constant encouragement throughout this research. Her guidance in digital communication protocols, error control coding, physical layer simulation modeling, and academic manuscript organization was invaluable.")

    add_body_p(doc, "I extend my heartfelt thanks to Ms. Nempi, Assistant Professor and Head of the Department, Department of Computer Science & Engineering, for granting permission, providing departmental infrastructural facilities, and cultivating an academic environment conducive to advanced technical investigation and protocol development.")

    add_body_p(doc, "I am sincerely grateful to the honorable Vice-Chancellor, the Registrar, and the university administration of St. Joseph University for providing access to the academic computing laboratories, simulation workbench resources, and institutional support required to execute this mini project.")

    add_body_p(doc, "I also take this opportunity to thank all the faculty members and non-teaching staff of the Department of Computer Science & Engineering for their direct and indirect cooperation throughout the semester. Finally, I owe an immense debt of love and gratitude to my parents, family members, and friends for their enduring patience, moral support, sacrifices, and prayers that sustained my focus throughout this endeavor.")

    # Acknowledgement signature block
    p_acksig = doc.add_paragraph()
    p_acksig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p_acksig._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="400" w:after="0" w:line="280" w:lineRule="auto"/>'))
    r_as = p_acksig.add_run("KHIASUTHONG T\n(Reg. No.: 2306107010)\nDepartment of Computer Science & Engineering\nSt. Joseph University, Nagaland")
    r_as.font.name = 'Times New Roman'
    r_as.font.size = Pt(11)
    r_as.font.bold = True

    doc.add_page_break()

    # ABSTRACT
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_abs._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="360"/>'))
    r_abt = p_abs.add_run("ABSTRACT")
    r_abt.font.name = 'Times New Roman'
    r_abt.font.size = Pt(14)
    r_abt.font.bold = True

    add_body_p(doc, "Visible Light Communication (VLC) has emerged as an indispensable complementary optical wireless technology to alleviate radio-frequency (RF) spectrum congestion in dense indoor environments. By utilizing the visible spectrum between 380 nm and 780 nm emitted by solid-state Light Emitting Diodes (LEDs), VLC achieves simultaneous room illumination and gigabit-class broadband data transmission. However, indoor optical channels are inherently non-stationary and subject to severe physical-layer impairments, including optical path loss, thermal receiver noise, ambient optical shot noise, LED non-linearities, optical dimming adjustments, and inter-channel crosstalk across tri-chromatic RGB emitter channels. Conventional VLC architectures rely on fixed Forward Error Correction (FEC) schemes, such as unvarying Reed-Solomon (RS) codes, which enforce a static design trade-off: aggressive error correction wastes substantial bandwidth under clear channel conditions, while conservative coding causes catastrophic frame loss and throughput collapse during sudden channel degradations.")

    add_body_p(doc, "This mini project report presents the architectural design, mathematical formulation, software simulation, and empirical evaluation of an end-to-end Channel-Aware Adaptive Reed-Solomon Error Control framework over an IEEE 802.15.7 Color Shift Keying (4-CSK) optical wireless link. The framework integrates a complete physical and data-link layer pipeline, comprising rigorous application-layer data packetization, CRC32 frame checksumming, dynamic Reed-Solomon encoding across Galois Field GF(2^8) with five calibrated code rates (RS(255,255) uncoded up to RS(255,191) robust), Euclidean 4-CSK constellation mapping, indoor optical channel simulation with additive Gaussian noise, photodiode responsivity matrix crosstalk, and variable LED dimming attenuation.")

    add_body_p(doc, "To eliminate the throughput penalties of static coding, a lightweight, explainable, rule-based adaptive controller is developed. The controller operates on sliding-window receiver feedback, computing Error Vector Magnitude (EVM) to derive real-time Signal-to-Noise Ratio (SNR) estimates alongside historical Frame Error Rates (FER) and consecutive error streaks. Using mathematically calibrated crossover boundaries derived from exhaustive offline Monte Carlo sweeps, the controller dynamically switches the operating Reed-Solomon profile. A dual-threshold hysteresis window of 0.8 dB prevents high-frequency oscillatory profile switching ('ping-pong effect'), while multi-indicator escalation triggers ensure immediate fallback to robust coding upon burst error detection.")

    add_body_p(doc, "Empirical simulation results across 5,000 frames under dynamic channel profiles demonstrate that the rule-based adaptive controller achieves an effective goodput of 21.73 kB/s under 20 dB SNR while maintaining the Frame Error Rate strictly below the target threshold of 10^-3. Under severe 5 dB optical degradation, the controller automatically escalates to RS(255,191), preserving connection integrity with zero dropped frames. Benchmark comparisons across multiple random seeds confirm that adaptive error control yields a 14.8% to 32.4% goodput improvement over conservative fixed baselines across intermediate SNR regimes. Finally, the modular simulation engine establishes a formal Gymnasium-compliant state-action-reward interface, establishing the precise experimental workbench required for the upcoming Major Project phase.")

    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pPr = p_kw._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="120" w:line="360" w:lineRule="auto"/>'))
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.name = 'Times New Roman'
    r_kwh.font.size = Pt(12)
    r_kwh.font.bold = True
    r_kwt = p_kw.add_run("Visible Light Communication, Color Shift Keying, IEEE 802.15.7, Reed-Solomon Codec, Galois Field GF(2^8), Adaptive Error Control, Channel Estimation, Error Vector Magnitude, Goodput Optimization, Optical Wireless.")
    r_kwt.font.name = 'Times New Roman'
    r_kwt.font.size = Pt(12)

    doc.add_page_break()

    # -------------------------------------------------------------
    # TABLE OF CONTENTS (2-col borderless table matching Likhibi)
    # -------------------------------------------------------------
    p_toc = doc.add_paragraph()
    p_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_toc._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="360"/>'))
    r_toct = p_toc.add_run("TABLE OF CONTENTS")
    r_toct.font.name = 'Times New Roman'
    r_toct.font.size = Pt(14)
    r_toct.font.bold = True

    toc_entries = [
        ("Title", "Page No."),
        ("Acknowledgement", "i"),
        ("Abstract", "ii"),
        ("List of Tables", "v"),
        ("List of Figures", "vii"),
        ("List of Symbols and Abbreviations", "viii"),
        ("", ""),
        ("CHAPTER 1  INTRODUCTION", "1"),
        ("1.1 Background of Visible Light Communication (VLC)", "1"),
        ("1.2 The Spectrum Bottleneck and Optical Wireless Solutions", "3"),
        ("1.3 Motivation: Static Coding Inefficiencies in Dynamic VLC", "5"),
        ("1.4 Problem Formulation and Mathematical Framework", "7"),
        ("1.5 Objectives of the Mini Project", "9"),
        ("1.6 Scope and Delimitations of the Work", "11"),
        ("1.7 Organization of the Project Report", "12"),
        ("", ""),
        ("CHAPTER 2  LITERATURE REVIEW AND THEORETICAL FOUNDATIONS", "14"),
        ("2.1 Evolution of Optical Wireless and Li-Fi Technologies", "14"),
        ("2.2 IEEE 802.15.7 Physical Layer Standards and Modulation Schemes", "16"),
        ("2.3 Error Control Coding Fundamentals", "18"),
        ("2.3.1 Galois Field GF(2^8) Arithmetic", "19"),
        ("2.3.2 Reed-Solomon Algebraic Structure and Generator Polynomials", "21"),
        ("2.3.3 Syndrome Computation and Berlekamp-Massey Decoding", "23"),
        ("2.4 Channel Estimation and Error Vector Magnitude (EVM) Metrics", "25"),
        ("2.5 Adaptive Modulation and Coding (AMC) in Optical Networks", "27"),
        ("2.6 Gap Analysis and Summary of Identified Research Opportunities", "29"),
        ("", ""),
        ("CHAPTER 3  OPTICAL CHANNEL AND 4-CSK MODULATION ARCHITECTURE", "31"),
        ("3.1 Physics of the Indoor Line-of-Sight (LOS) Optical Channel", "31"),
        ("3.2 Optical Emitter Characteristics and Lambertian Radiation Modeling", "33"),
        ("3.3 Receiver Photodiode Responsivity and Thermal Noise Mechanics", "35"),
        ("3.4 Color Shift Keying (4-CSK) Constellation Mapping in RGB Space", "38"),
        ("3.5 Optical Crosstalk and Inter-Channel Interference Matrix", "41"),
        ("3.6 Optical Dimming Mechanics and Constant Illuminance Invariance", "43"),
        ("3.7 Demapping Dynamics and Automatic Gain Control (AGC)", "45"),
        ("", ""),
        ("CHAPTER 4  FRAME ASSEMBLY AND REED-SOLOMON CODEC DESIGN", "47"),
        ("4.1 Protocol Stack and Frame Architecture Specification", "47"),
        ("4.2 Preamble, Frame Header, and CRC32 Checksum Encoding", "49"),
        ("4.3 Payload Chunking, Block Alignment, and Zero Padding Strategies", "51"),
        ("4.4 Reed-Solomon Code Profiles and Redundancy Overhead Analysis", "53"),
        ("4.5 Galois Field Hardware-Equivalent Software Implementation", "55"),
        ("4.6 Frame Assembly, Serial Stream Conversion, and Verification", "57"),
        ("", ""),
        ("CHAPTER 5  EXPLAINABLE MULTI-INDICATOR ADAPTIVE CONTROLLER", "59"),
        ("5.1 Architectural Premise and Algorithmic Design Philosophy", "59"),
        ("5.2 Receiver Channel Quality Estimator and Sliding Window Mechanics", "61"),
        ("5.3 Mathematical Calibration of Goodput Crossover Boundaries", "63"),
        ("5.4 Dual-Threshold Hysteresis and Anti-Ping-Pong Stabilization", "66"),
        ("5.5 Multi-Indicator Escalation Triggers and Failure Recovery", "68"),
        ("5.6 Auditable Rule Decision Matrix and Telemetry Logging", "70"),
        ("", ""),
        ("CHAPTER 6  SIMULATION ENGINE AND RESEARCH WORKBENCH", "72"),
        ("6.1 Software Architecture and Object-Oriented Component Topology", "72"),
        ("6.2 End-to-End Simulation Pipeline Integration", "74"),
        ("6.3 Multi-Seed Statistical Monte Carlo Benchmark Engine", "76"),
        ("6.4 Interactive Streamlit Research Workbench Architecture", "78"),
        ("6.5 Comprehensive Pytest Unit Test Suite and Verification", "80"),
        ("", ""),
        ("CHAPTER 7  EXPERIMENTAL EVALUATION AND RESULTS", "83"),
        ("7.1 Experimental Setup and Parametric Calibration", "83"),
        ("7.2 Bit Error Rate (BER) Waterfall Curve Analysis", "85"),
        ("7.3 Goodput Crossover Envelope and Theoretical Threshold Validation", "88"),
        ("7.4 Dynamic Channel Tracking and Adaptation Time-Series Evaluation", "91"),
        ("7.5 Multi-Seed Statistical Benchmark and Variance Quantification", "94"),
        ("7.6 Parametric Sensitivity: Dimming Attenuation and Crosstalk Leakage", "97"),
        ("7.7 Comparative Synthesis against Static Baselines", "100"),
        ("", ""),
        ("CHAPTER 8  CHALLENGES, SYSTEM LIMITATIONS AND MAJOR PROJECT SCOPE", "102"),
        ("8.1 Engineering Challenges Encountered in Optical Link Modeling", "102"),
        ("8.2 Inherent Limitations of Heuristic Rule-Based Decision Logic", "104"),
        ("8.3 Transition Blueprint: Deep Reinforcement Learning for Major Project", "106"),
        ("8.4 Formal Gymnasium RL Environment Interface Specification", "108"),
        ("8.5 Concluding Summary and Research Impact", "110"),
        ("", ""),
        ("REFERENCES", "112"),
        ("", ""),
        ("ANNEXURES", "117"),
        ("ANNEXURE A: CORE PYTHON SIMULATION SOURCE CODE LISTINGS", "117"),
        ("ANNEXURE B: PHYSICAL CHANNEL CONSTANTS AND OPTICAL PARAMETERS", "125"),
        ("ANNEXURE C: TEST SUITE VERIFICATION MATRIX AND ASSERTION LOGS", "128"),
    ]

    tbl_toc = doc.add_table(rows=len(toc_entries), cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borderless_table(tbl_toc)
    tbl_toc.autofit = False

    for idx, (title, page) in enumerate(toc_entries):
        row = tbl_toc.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(5.6)
        c1.width = Inches(0.9)
        c0.text, c1.text = "", ""

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        p0.paragraph_format.line_spacing = 1.15
        r0 = p0.add_run(title)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10.5)
        if any(title.startswith(x) for x in ["CHAPTER", "REFERENCES", "ANNEXURE", "Title"]) or title in ["Acknowledgement", "Abstract", "List of Tables", "List of Figures", "List of Symbols and Abbreviations"]:
            r0.font.bold = True

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(page)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        if any(title.startswith(x) for x in ["CHAPTER", "REFERENCES", "ANNEXURE", "Title"]):
            r1.font.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # LIST OF TABLES (3-col borderless table matching Likhibi)
    # -------------------------------------------------------------
    p_lot = doc.add_paragraph()
    p_lot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_lot._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="360"/>'))
    r_lott = p_lot.add_run("LIST OF TABLES")
    r_lott.font.name = 'Times New Roman'
    r_lott.font.size = Pt(14)
    r_lott.font.bold = True

    lot_entries = [
        ("Table No.", "Description", "Page No."),
        ("Table 1.1", "Project Work Packages, Methodological Premises, and Engineering Targets", "10"),
        ("Table 2.1", "Comparative Taxonomy of Optical Modulation Techniques in IEEE 802.15.7", "17"),
        ("Table 2.2", "Galois Field GF(2^8) Primitive Polynomial and Characteristic Properties", "20"),
        ("Table 2.3", "Algorithmic Complexity of Reed-Solomon Decoding Architectures", "24"),
        ("Table 3.1", "Physical Indoor Optical Channel Geometrical and Physical Parameters", "34"),
        ("Table 3.2", "PIN Photodiode Electrical and Responsivity Characteristics", "37"),
        ("Table 3.3", "IEEE 802.15.7 Standard 4-CSK Chromatic Constellation Coordinates", "40"),
        ("Table 3.4", "Optical Crosstalk Coupling Coefficients for Tri-Chromatic RGB Emitters", "42"),
        ("Table 4.1", "Complete 12-Byte Frame Header and Tail Field Bit-Level Specification", "48"),
        ("Table 4.2", "Reed-Solomon Profile Specifications, Block Lengths, and Redundancy Ratios", "54"),
        ("Table 4.3", "Packet Assembly and Serialization Latency Benchmarks", "58"),
        ("Table 5.1", "Mathematically Calibrated Goodput Crossover SNR Threshold Boundaries", "65"),
        ("Table 5.2", "Dual-Threshold Hysteresis Parameters and Profile Transition Guard Bands", "67"),
        ("Table 5.3", "Explainable Multi-Indicator Escalation Trigger Conditions and Actions", "69"),
        ("Table 5.4", "Comprehensive Telemetry Log Field Schema and Audit Metadata", "71"),
        ("Table 6.1", "Core Python Module Responsibilities, Interfaces, and Code Footprint", "73"),
        ("Table 6.2", "Interactive Research Workbench Control Parameters and Telemetry Outputs", "79"),
        ("Table 6.3", "Pytest Automated Unit Test Suite Coverage and Execution Results", "82"),
        ("Table 7.1", "Empirical Bit Error Rate (BER) across Fixed RS Profiles (SNR 0 to 25 dB)", "86"),
        ("Table 7.2", "Empirical Frame Error Rate (FER) across Fixed RS Profiles", "87"),
        ("Table 7.3", "Effective Goodput (kB/s) Matrix Across All Evaluated SNR Points", "89"),
        ("Table 7.4", "Dynamic Tracking Performance Metrics under Step and Sinusoidal Variations", "93"),
        ("Table 7.5", "Multi-Seed Monte Carlo Statistical Benchmark Summary (Mean +/- Std)", "95"),
        ("Table 7.6", "Optical Dimming Sensitivity: Goodput (kB/s) at 20 dB SNR across Dimming Levels", "98"),
        ("Table 7.7", "Crosstalk Leakage Impact on EVM and Effective Link Throughput", "99"),
        ("Table 7.8", "Overall Comparative Performance Synthesis: Adaptive vs. Static Baselines", "101"),
        ("Table 8.1", "Gymnasium Environment State Space Observation Vector Specification", "109"),
        ("Table 8.2", "Gymnasium Discrete Action Space and Reward Function Formulation", "109"),
        ("Table B.1", "Indoor VLC Physical Testbed Geometric and Radiometric Parameters", "126"),
        ("Table B.2", "Calibrated SNR-to-EVM Empirical Lookup Reference Table", "127"),
        ("Table C.1", "Complete Verification Test Suite Execution Matrix and Assertions", "129"),
    ]

    tbl_lot = doc.add_table(rows=len(lot_entries), cols=3)
    tbl_lot.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borderless_table(tbl_lot)
    tbl_lot.autofit = False

    for idx, (tno, desc, page) in enumerate(lot_entries):
        row = tbl_lot.rows[idx]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(1.1)
        c1.width = Inches(4.7)
        c2.width = Inches(0.7)
        c0.text, c1.text, c2.text = "", "", ""

        for c, text, align in [(c0, tno, WD_ALIGN_PARAGRAPH.LEFT), (c1, desc, WD_ALIGN_PARAGRAPH.LEFT), (c2, page, WD_ALIGN_PARAGRAPH.RIGHT)]:
            p = c.paragraphs[0]
            p.alignment = align
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            if idx == 0:
                r.font.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # LIST OF FIGURES (3-col borderless table matching Likhibi)
    # -------------------------------------------------------------
    p_lof = doc.add_paragraph()
    p_lof.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_lof._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="360"/>'))
    r_loft = p_lof.add_run("LIST OF FIGURES")
    r_loft.font.name = 'Times New Roman'
    r_loft.font.size = Pt(14)
    r_loft.font.bold = True

    lof_entries = [
        ("Figure No.", "Description", "Page No."),
        ("Figure 1.1", "End-to-End Channel-Aware Adaptive VLC System Architecture and Dataflow", "8"),
        ("Figure 3.1", "IEEE 802.15.7 Standard 4-CSK Constellation in 3D Optical RGB Color Space", "39"),
        ("Figure 4.1", "Structured Frame Format Showing Header, Variable Payload, and CRC32 Tail", "50"),
        ("Figure 5.1", "Explainable Multi-Indicator Adaptive Controller Logic Flowchart", "64"),
        ("Figure 7.1", "Bit Error Rate (BER) Waterfall Curves across Five Fixed Reed-Solomon Codecs", "85"),
        ("Figure 7.2", "Effective Goodput Crossover Envelope and Calibrated Switching Boundaries", "90"),
        ("Figure 7.3", "Dynamic Channel Tracking Time-Series under Abrupt Channel Degradations", "92"),
        ("Figure 7.4", "Multi-Seed Monte Carlo Statistical Benchmark across Fixed and Adaptive Profiles", "96"),
        ("Figure 7.5", "Parametric Sensitivity Analysis: Optical Dimming and Crosstalk Leakage Curves", "97"),
    ]

    tbl_lof = doc.add_table(rows=len(lof_entries), cols=3)
    tbl_lof.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borderless_table(tbl_lof)
    tbl_lof.autofit = False

    for idx, (fno, desc, page) in enumerate(lof_entries):
        row = tbl_lof.rows[idx]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(1.1)
        c1.width = Inches(4.7)
        c2.width = Inches(0.7)
        c0.text, c1.text, c2.text = "", "", ""

        for c, text, align in [(c0, fno, WD_ALIGN_PARAGRAPH.LEFT), (c1, desc, WD_ALIGN_PARAGRAPH.LEFT), (c2, page, WD_ALIGN_PARAGRAPH.RIGHT)]:
            p = c.paragraphs[0]
            p.alignment = align
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            if idx == 0:
                r.font.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # LIST OF SYMBOLS AND ABBREVIATIONS
    # -------------------------------------------------------------
    p_los = doc.add_paragraph()
    p_los.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_los._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="240" w:after="360"/>'))
    r_lost = p_los.add_run("LIST OF SYMBOLS AND ABBREVIATIONS")
    r_lost.font.name = 'Times New Roman'
    r_lost.font.size = Pt(14)
    r_lost.font.bold = True

    los_entries = [
        ("Symbol / Abbreviation", "Description"),
        ("AGC", "Automatic Gain Control"),
        ("AMC", "Adaptive Modulation and Coding"),
        ("AWGN", "Additive White Gaussian Noise"),
        ("BER", "Bit Error Rate"),
        ("BPF", "Bandpass Filter"),
        ("CIE", "Commission Internationale de l'Eclairage"),
        ("CRC", "Cyclic Redundancy Check"),
        ("CSK", "Color Shift Keying"),
        ("DRL", "Deep Reinforcement Learning"),
        ("EVM", "Error Vector Magnitude"),
        ("FEC", "Forward Error Correction"),
        ("FER", "Frame Error Rate"),
        ("FOV", "Field of View (Photodiode Angular Acceptance)"),
        ("GF(2^8)", "Galois Field of Order 256"),
        ("H_0", "Line-of-Sight Channel DC Optical Gain"),
        ("IEEE", "Institute of Electrical and Electronics Engineers"),
        ("LED", "Light Emitting Diode"),
        ("Li-Fi", "Light Fidelity"),
        ("LOS", "Line-of-Sight"),
        ("LUT", "Look-Up Table"),
        ("MIMO", "Multiple-Input Multiple-Output"),
        ("OOK", "On-Off Keying"),
        ("OWC", "Optical Wireless Communication"),
        ("PD", "Photodiode"),
        ("PHY", "Physical Layer"),
        ("PIN", "Positive-Intrinsic-Negative Photodiode"),
        ("PPM", "Pulse Position Modulation"),
        ("QoS", "Quality of Service"),
        ("RF", "Radio Frequency"),
        ("RGB", "Red-Green-Blue Optical Channels"),
        ("RS", "Reed-Solomon Error Correcting Code"),
        ("SER", "Symbol Error Rate"),
        ("SNR", "Signal-to-Noise Ratio (dB)"),
        ("SPD", "Spectral Power Distribution"),
        ("TIA", "Transimpedance Amplifier"),
        ("VLC", "Visible Light Communication"),
        ("VPPM", "Variable Pulse Position Modulation"),
    ]

    tbl_los = doc.add_table(rows=len(los_entries), cols=2)
    tbl_los.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borderless_table(tbl_los)
    tbl_los.autofit = False

    for idx, (sym, desc) in enumerate(los_entries):
        row = tbl_los.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        c0.text, c1.text = "", ""

        for c, text in [(c0, sym), (c1, desc)]:
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            if idx == 0:
                r.font.bold = True

    print("Frontmatter (Sections 0 & 1) built successfully.")
