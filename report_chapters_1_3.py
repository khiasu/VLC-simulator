import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_helpers import set_cell_margins, set_cell_shading, set_table_borders, make_equation_box

def add_chapters_1_to_3(doc, add_p, add_heading_1, add_heading_2, add_chapter_title, add_figure_with_caption):
    # =========================================================================
    # CHAPTER 1: INTRODUCTION
    # =========================================================================
    add_chapter_title(1, "Introduction")
    
    add_heading_1("1.1 Motivation and Optical Wireless Communication Background")
    add_p(
        "Over the past decade, the exponential proliferation of wireless personal devices, high-definition multimedia "
        "streaming, and Internet of Things (IoT) hardware has placed unprecedented demands on conventional radio frequency "
        "(RF) communication networks. The existing RF spectrum—predominantly spanning the sub-6 GHz industrial, scientific, "
        "and medical (ISM) bands—is experiencing severe spectral congestion. In dense indoor venues such as corporate auditoriums, "
        "commercial passenger aircraft, healthcare facilities, and automated industrial plants, RF systems frequently suffer from "
        "co-channel interference, restricted channel bandwidths, and strict regulatory limitations concerning electromagnetic "
        "radiation emissions."
    )
    add_p(
        "Visible Light Communication (VLC), a specialized subset of Optical Wireless Communication (OWC), provides an effective "
        "complementary alternative by utilizing the visible optical spectrum (spanning approximately 380 nm to 780 nm, or 400 THz "
        "to 790 THz). By leveraging existing solid-state lighting infrastructures—specifically multi-color Light Emitting Diodes "
        "(LEDs)—VLC systems achieve dual functionality: indoor space illumination and concurrent gigabit-class wireless data "
        "telemetry. Because visible optical photons do not pass through opaque architectural walls, VLC offers inherent physical "
        "layer confinement, significantly curtailing malicious eavesdropping without the processing overhead of high-complexity "
        "cryptographic primitives. Furthermore, visible optical wavelengths cause zero electromagnetic interference, permitting "
        "safe operation in environments where RF transmissions are prohibited, such as hospital intensive care units and explosive "
        "petrochemical manufacturing bays."
    )
    add_p(
        "Despite these substantial benefits, indoor optical wireless channels are susceptible to dynamic environmental "
        "impairments. In practical indoor deployments, the instantaneous channel quality fluctuates over time due to three "
        "primary physical phenomena: user mobility and receiver orientation changes relative to the ceiling luminaire, ambient "
        "illumination interference from direct sunlight and artificial fluorescent fixtures, and user-adjusted luminous dimming. "
        "Consequently, static physical and link-layer designs perform suboptimally under dynamic indoor optical conditions."
    )

    add_heading_1("1.2 Problem Formulation and Mathematical Framework")
    add_p(
        "Traditional digital communication architectures commonly employ fixed Forward Error Correction (FEC) codes designed "
        "for the anticipated worst-case channel Signal-to-Noise Ratio (SNR). In optical wireless networks, this static approach "
        "introduces a critical operational inefficiency:"
    )
    add_p(
        "1. Heavy Parity Overhead in Benign Channels: When the optical receiver is located directly underneath the transmitting "
        "luminaire under optimal Line-of-Sight (LOS) conditions (yielding channel SNRs exceeding 20 dB), a static heavy code—such "
        "as a Reed-Solomon RS(255, 223) code requiring 32 parity bytes per 255-byte block—wastes approximately 12.5% of the "
        "available optical bandwidth on redundant parity symbols that are rarely required."
    )
    add_p(
        "2. Link Outage in Degraded Channels: Conversely, if a lightweight code—such as RS(255, 251) with only 4 parity bytes—or "
        "uncoded transmission is selected to maximize throughput, any transient link attenuation (caused by receiver displacement, "
        "illumination dimming, or ambient noise surges) causes the channel error rate to exceed the error-correction bound. This "
        "results in uncorrectable frame corruptions and severe packet drop rates."
    )
    add_p(
        "To resolve this trade-off, this mini project formulates an adaptive cross-layer optimization problem. Let S denote the "
        "set of available Reed-Solomon coding profiles over Galois Field GF(2^8):"
    )
    make_equation_box(doc, "S = { RS(255, 223), RS(255, 239), RS(255, 247), RS(255, 251), UNCODED }", "(1.1)")
    add_p(
        "At each transmission frame index t, the primary objective is to select an optimal coding profile s_t in S that "
        "maximizes the useful application Goodput G(s_t, SNR_t) while satisfying a strict link reliability constraint on the "
        "Frame Error Rate (FER):"
    )
    make_equation_box(doc, "s_t^* = argmax_{s in S} G(s, SNR_t)    subject to:  FER(s, SNR_t) <= 10^{-2}", "(1.2)")
    add_p(
        "where Goodput is defined mathematically as the rate of error-free, useful payload bits successfully delivered to the "
        "receiving application layer per unit time:"
    )
    make_equation_box(doc, "G(s, SNR) = [ L_payload * 8 * (1 - FER(s, SNR)) ] / T_frame    (kbps)", "(1.3)")
    add_p(
        "Here, L_payload represents the unpadded payload length in bytes, and T_frame denotes the total physical transmission "
        "duration of the frame over the optical channel."
    )

    add_heading_1("1.3 Physical and Data Link Layer Context of Visible Light Communication")
    add_p(
        "To implement dynamic error control effectively, the simulation framework models the interaction between the Data Link "
        "Layer and the Physical Layer defined by the IEEE 802.15.7 standard. While traditional RF systems modulate signal phase "
        "and amplitude, optical communication uses Intensity Modulation with Direct Detection (IM/DD). Among the optical modulation "
        "schemes standardized in IEEE 802.15.7, Color Shift Keying (CSK) provides distinct advantages by modulating data into the "
        "relative optical intensities of multi-color (Red, Green, Blue) LEDs while keeping the total instantaneous radiant flux "
        "constant (R + G + B = 1.0). This property eliminates perceptible optical flicker to human occupants."
    )
    add_p(
        "Figure 1.1 illustrates the complete layered communication architecture designed and evaluated in this project. The "
        "pipeline bridges raw bitstream generation, structured framing with cyclic redundancy verification, Galois field block "
        "encoding, optical 4-CSK symbol modulation, optical channel impairment injection, and real-time receiver feedback."
    )
    add_figure_with_caption("figures/fig_pipeline_architecture.png", "Figure 1.1: End-to-End Layered Communication Architecture of the Adaptive VLC Simulator", width_in=6.0)

    add_heading_1("1.4 Objectives of the Mini Project")
    add_p("The specific engineering and research objectives established for this mini project are:")
    add_p(
        "1. Modular Communication Pipeline: Construct an end-to-end discrete-event software simulation platform in Python modeling "
        "the physical and data link layers of an indoor VLC link with full parameter traceability."
    )
    add_p(
        "2. Dynamic Galois Field Codec: Implement a dynamic block-chunked Reed-Solomon encoder and decoder over GF(2^8) capable of "
        "on-the-fly reconfiguration across multiple parity strengths without restarting execution state."
    )
    add_p(
        "3. Physical Optical Channel Modeling: Implement IEEE 802.15.7 4-CSK modulation with constant total power and model "
        "realistic indoor optical impairments, including additive Gaussian noise, LED dimming scaling, and inter-channel RGB spectral crosstalk."
    )
    add_p(
        "4. Real-Time Receiver Estimation: Design a receiver channel quality estimator that converts Error Vector Magnitude (EVM) "
        "from received 4-CSK symbols into estimated channel SNR and maintains sliding-window link health statistics."
    )
    add_p(
        "5. Explainable Adaptive Controller: Develop a deterministic, rule-based adaptive controller incorporating data-driven "
        "crossover thresholds, a 0.5 dB hysteresis margin to prevent ping-pong oscillations, and an emergency escalation fallback."
    )
    add_p(
        "6. Empirical Benchmarking and Workbench: Conduct multi-seed dynamic benchmark experiments, generate publication-grade "
        "analytical plots, and construct an interactive web-based research dashboard for real-time telemetry and frame inspection."
    )

    add_heading_1("1.5 Scope and Delimitations of the Mini Project")
    add_p(
        "The scope of this mini project is specifically focused on establishing an explainable, deterministic rule-based adaptive "
        "control mechanism and validating the complete physical and data link software pipeline. To ensure rigorous evaluation within "
        "the minor project framework, the following delimitations apply:"
    )
    add_p(
        "1. Rule-Based vs. AI Control: In accordance with academic guidelines for this phase, the adaptive controller uses deterministic "
        "threshold heuristics and sliding-window statistical indicators. Complex artificial intelligence and Deep Reinforcement "
        "Learning (DRL) algorithms (such as PPO or DQN) are excluded from this report and reserved for the major project phase, for which "
        "this verified simulator serves as the foundational Gymnasium environment."
    )
    add_p(
        "2. Software Simulation: The investigation is conducted via numerical software simulation with realistic physical parameters "
        "derived from empirical optical literature. Hardware photodiode front-ends and physical driver circuits are outside the scope "
        "of this mini project."
    )

    add_heading_1("1.6 Organization of the Project Report")
    add_p(
        "This project report is organized into eight chapters. Chapter 2 presents a literature review covering VLC channel models, "
        "CSK modulation, Reed-Solomon coding theory over GF(2^8), and existing optical simulation tools. Chapter 3 details the optical "
        "channel model, 4-CSK modulation geometry, and physical impairment formulations. Chapter 4 describes the data link framing protocol, "
        "CRC32 integrity checks, and the dynamic RS codec architecture. Chapter 5 formulates the receiver quality estimator and the "
        "multi-indicator adaptive controller. Chapter 6 outlines the simulation engine, interactive research dashboard, and pytest "
        "verification suite. Chapter 7 presents empirical results, crossover boundary derivations, multi-seed benchmarks, and sensitivity "
        "analyses. Finally, Chapter 8 concludes the report and discusses the future scope of continuation for the major project."
    )

    # =========================================================================
    # CHAPTER 2: LITERATURE REVIEW AND THEORETICAL FOUNDATIONS
    # =========================================================================
    add_chapter_title(2, "Literature Review and Theoretical Foundations")
    
    add_heading_1("2.1 Optical Wireless Communication and VLC Channel Models")
    add_p(
        "Optical Wireless Communication has evolved from early infrared links into high-speed visible light networking. In indoor "
        "VLC environments, optical propagation is governed by radiant flux emissions from phosphor-converted white LEDs or multi-chip "
        "RGB LED arrays. The optical transmitter follows a generalized Lambertian radiation pattern, where the radiant intensity "
        "R_o(phi) depends on the emission angle phi:"
    )
    make_equation_box(doc, "R_o(phi) = [ (m + 1) / (2 * pi) ] * cos^m(phi)", "(2.1)")
    add_p(
        "where m is the Lambertian order, determined by the transmitter semi-angle at half power Phi_{1/2} through the relation "
        "m = -ln(2) / ln(cos(Phi_{1/2})). At the receiver side, a silicon p-i-n photodiode or avalanche photodiode (APD) captures the "
        "incident photons, producing a proportional photocurrent. The direct Line-of-Sight (LOS) optical channel direct-current (DC) gain "
        "H_{LOS}(0) between transmitter and receiver separated by physical distance d is expressed as:"
    )
    make_equation_box(doc, "H_{LOS}(0) = [ (m + 1) * A_det / (2 * pi * d^2) ] * cos^m(phi) * T_s(psi) * g(psi) * cos(psi)", "(2.2)")
    add_p(
        "where A_det is the physical detection area of the photodiode, psi is the angle of incidence relative to the receiver surface "
        "normal, T_s(psi) is the transmission coefficient of the optical bandpass filter, and g(psi) is the internal optical concentrator gain:"
    )
    make_equation_box(doc, "g(psi) = n_c^2 / sin^2(Psi_c)    for  0 <= psi <= Psi_c", "(2.3)")
    add_p(
        "where n_c represents the refractive index of the concentrator medium and Psi_c denotes the internal field-of-view (FOV) semi-angle. "
        "When the incidence angle psi exceeds Psi_c, the received optical power drops to zero."
    )

    add_heading_1("2.2 Color Shift Keying (CSK) Modulation Theory")
    add_p(
        "The IEEE 802.15.7 standard defines Color Shift Keying (CSK) for VLC systems utilizing tri-color RGB light sources. Unlike "
        "conventional On-Off Keying (OOK) or Pulse Position Modulation (PPM), which vary LED operational duty cycles and introduce "
        "risk of observable human eye flicker during dimming, CSK operates by modulating the emitted chromaticity coordinates in the "
        "standard CIE 1931 color space while constraining total emitted instantaneous luminous power to a constant value."
    )
    add_p(
        "A CSK constellation consists of M discrete points in the CIE 1931 chromaticity diagram, where each point is characterized "
        "by chromaticity coordinates (x_k, y_k). By applying color conversion equations, these coordinates map to instantaneous "
        "normalized optical drive currents for the Red, Green, and Blue emitters:"
    )
    make_equation_box(doc, "I_R + I_G + I_B = 1.0,    with  I_R, I_G, I_B >= 0", "(2.4)")
    add_p(
        "This constant optical power constraint ensures that human observers perceive only steady white or tinted ambient room illumination, "
        "even when the individual color channels are modulated at megahertz switching speeds."
    )

    add_heading_1("2.3 Forward Error Correction Fundamentals (Reed-Solomon Codes)")
    add_p(
        "Reed-Solomon (RS) codes are non-binary linear cyclic block codes introduced by Irving S. Reed and Gustave Solomon in 1960. "
        "Operating over Galois Fields GF(2^m), RS codes exhibit optimal Minimum Distance Separable (MDS) properties, meaning they "
        "achieve the maximum possible error-correcting capability for a given parity check overhead."
    )

    add_heading_2("2.3.1 Galois Field Arithmetic and Generator Polynomials")
    add_p(
        "In this simulator, all Reed-Solomon operations are defined over the finite field GF(2^8) (containing 256 discrete elements), "
        "matching the standard 8-bit byte representation of computer architectures. The field GF(2^8) is constructed modulo the primitive "
        "polynomial p(x):"
    )
    make_equation_box(doc, "p(x) = x^8 + x^4 + x^3 + x^2 + 1", "(2.5)")
    add_p(
        "Let alpha be a primitive root of p(x). Every non-zero field element can be expressed as a distinct power alpha^i for 0 <= i <= 254. "
        "Addition in GF(2^8) corresponds to bitwise XOR, while multiplication is performed modulo p(x). An RS(N, K) code takes K information "
        "symbols and appends 2t = N - K parity symbols to form an N-symbol codeword polynomial c(x). The generator polynomial g(x) is given by:"
    )
    make_equation_box(doc, "g(x) = prod_{i=0}^{2t-1} (x - alpha^{b+i})", "(2.6)")
    add_p("where b is an integer offset (typically chosen as b = 0 or b = 1).")

    add_heading_2("2.3.2 Parity Redundancy and Error-Correction Capability")
    add_p(
        "The minimum Hamming distance of an MDS code satisfies d_min = N - K + 1. Consequently, an RS(N, K) code can correct up to "
        "t symbol (byte) errors occurring at unknown locations within the block, where:"
    )
    make_equation_box(doc, "t = floor( (N - K) / 2 )", "(2.7)")
    add_p(
        "Because errors are evaluated on a symbol basis rather than bit by bit, any arbitrary bit corruption (from 1 to 8 bits) within "
        "a single byte consumes only one unit of correction capacity t. This makes Reed-Solomon codes resilient against burst errors "
        "induced by optical fading and transient channel noise."
    )

    add_heading_2("2.3.3 Syndrome Computation and Berlekamp-Massey Decoding")
    add_p(
        "At the receiving end, the decoder evaluates the received polynomial r(x) = c(x) + e(x) at the roots of g(x) to compute "
        "2t syndrome values S_i = r(alpha^{b+i}) for 0 <= i < 2t. If all syndromes evaluate to zero, no detectable errors exist. "
        "If non-zero syndromes are detected, the Berlekamp-Massey algorithm determines the error locator polynomial sigma(x). "
        "The Chien search identifies the error locations X_k, and Forney's algorithm calculates the numerical error values Y_k, "
        "recovering the original codeword."
    )

    add_heading_1("2.4 Channel Quality Estimation Techniques (EVM to SNR Conversion)")
    add_p(
        "Real-time link adaptation requires low-latency channel quality feedback at the receiver. Traditional wireless systems rely "
        "on periodic pilot preambles, which consume transmission bandwidth. In this project, channel quality is estimated continuously "
        "from payload symbols using Error Vector Magnitude (EVM). EVM measures the root-mean-square displacement between the received "
        "noisy constellation vectors I_rx,i and their corresponding nearest ideal reference constellation points I_ref,i:"
    )
    make_equation_box(doc, "EVM = sqrt( [ sum_{i=1}^M || I_{rx,i} - I_{ref,i} ||^2 ] / [ sum_{i=1}^M || I_{ref,i} ||^2 ] )", "(2.8)")
    add_p(
        "Assuming zero-mean Gaussian channel noise, EVM relates inversely to the physical Signal-to-Noise Ratio. In decibel scale, the "
        "effective channel SNR is estimated as:"
    )
    make_equation_box(doc, "SNR_{est_dB} approx -20 * log_{10}(EVM)", "(2.9)")
    add_p(
        "This estimation approach eliminates the need for dedicated pilot symbol overhead, allowing passive channel state estimation directly "
        "from received 4-CSK optical frames."
    )

    add_heading_1("2.5 Adaptive Modulation and Coding (AMC) in Optical Networks")
    add_p(
        "Adaptive Modulation and Coding (AMC) dynamically matches physical and link layer parameters to instantaneous channel quality. "
        "In RF communications, AMC is widely adopted across standards such as LTE, 5G NR, and IEEE 802.11ax. However, direct porting "
        "of RF AMC schemes to indoor optical wireless channels is limited by key differences: unipolar optical power constraints, "
        "non-linear LED drive characteristics, and eye-safety/dimming requirements."
    )
    add_p(
        "A critical challenge in threshold-based link adaptation is preventing 'ping-pong' oscillations—frequent, unnecessary profile "
        "switching triggered by small noise fluctuations around decision boundaries. In this work, an explicit hysteresis margin "
        "Delta = 0.5 dB is integrated into the decision logic to ensure stable link state transitions under channel jitter."
    )

    add_heading_1("2.6 Comparative Summary of Existing Simulators and Research Gaps")
    add_p(
        "Table 2.1 summarizes the capabilities of existing optical wireless simulation tools compared with the framework developed in this project."
    )

    # Table 2.1
    comp_data = [
        ("Feature / Capability", "MATLAB VLC Toolbox", "NS-3 VLC Module", "Proposed Simulator"),
        ("Optical Modulation Support", "OOK / PPM only", "OOK only", "IEEE 802.15.7 4-CSK (RGB)"),
        ("Constant Optical Power Modeling", "No (Intensity swings)", "No (Single color)", "Yes (Sum I = 1.0 constraint)"),
        ("Dynamic Reed-Solomon Codec", "Static RS only", "None / Abstracted", "Dynamic GF(2^8) (5 Profiles)"),
        ("Multi-Impairment Optical Channel", "AWGN only", "Path loss only", "AWGN + Dimming + Crosstalk"),
        ("Passive EVM Channel Estimation", "No", "No", "Yes (Real-time EVM to SNR)"),
        ("Hysteresis & Escalation Controller", "No", "No", "Yes (0.5 dB margin + Burst trigger)"),
        ("Interactive Telemetry Dashboard", "None", "None", "Streamlit Workbench + Log Replay")
    ]
    t_comp = doc.add_table(rows=len(comp_data), cols=4)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_comp.autofit = False
    set_table_borders(t_comp, color="D0D0D0", sz="4")
    for r_i, row in enumerate(comp_data):
        for c_i, val in enumerate(row):
            cell = t_comp.rows[r_i].cells[c_i]
            cell.width = Inches(1.6) if c_i > 0 else Inches(2.2)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
            r = p.add_run(val); r.font.name = 'Times New Roman'; r.font.size = Pt(10)
            if r_i == 0:
                set_cell_shading(cell, "E8E8E8")
                r.font.bold = True
            elif c_i == 3:
                r.font.bold = True

    # =========================================================================
    # CHAPTER 3: OPTICAL CHANNEL MODELING AND 4-CSK PHYSICAL LAYER
    # =========================================================================
    add_chapter_title(3, "Optical Channel Modeling and 4-CSK Physical Layer")
    
    add_heading_1("3.1 Indoor Optical Propagation and Line-of-Sight Geometry")
    add_p(
        "Indoor optical wireless telemetry relies predominantly on the direct Line-of-Sight (LOS) path between the ceiling luminaire "
        "and the receiver photodetector. Diffuse multipath reflections off walls contribute slight intersymbol interference (ISI) at "
        "extremely high modulation bandwidths (exceeding 100 MHz), but for typical indoor symbol rates (1 to 10 MSymbols/s), the channel "
        "is predominantly non-frequency-selective (flat fading) with deterministic path loss determined by link distance and spatial alignment."
    )

    add_heading_1("3.2 IEEE 802.15.7 4-CSK Constellation Geometry and Power Normalization")
    add_p(
        "Color Shift Keying (CSK) modulates digital bitstreams into multi-color light emission vectors. Under the IEEE 802.15.7 4-CSK "
        "profile, every 2 input bits are mapped to one of four discrete optical intensity vectors in normalized three-dimensional "
        "Red-Green-Blue (RGB) color space. Table 3.1 details the bit-to-symbol allocation, color state assignments, and normalized optical powers."
    )

    # Table 3.1
    csk_map_data = [
        ("2-Bit Symbol", "Color State", "Red Intensity (I_R)", "Green Intensity (I_G)", "Blue Intensity (I_B)", "Total Power (Sum I)"),
        ("'00'", "Pure Red", "1.000", "0.000", "0.000", "1.000"),
        ("'01'", "Pure Green", "0.000", "1.000", "0.000", "1.000"),
        ("'10'", "Pure Blue", "0.000", "0.000", "1.000", "1.000"),
        ("'11'", "Neutral White", "0.333", "0.333", "0.333", "1.000")
    ]
    t_csk = doc.add_table(rows=len(csk_map_data), cols=6)
    t_csk.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_csk.autofit = False
    set_table_borders(t_csk, color="D0D0D0", sz="4")
    for r_i, row in enumerate(csk_map_data):
        for c_i, val in enumerate(row):
            cell = t_csk.rows[r_i].cells[c_i]
            cell.width = Inches(1.1)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
            r = p.add_run(val); r.font.name = 'Times New Roman'; r.font.size = Pt(10)
            if r_i == 0:
                set_cell_shading(cell, "E8E8E8")
                r.font.bold = True

    add_p(
        "Figure 3.1 illustrates the geometric distribution of the 4-CSK constellation points in normalized RGB space. The triangular "
        "plane represents the constant optical power boundary (R + G + B = 1.0). Transmitted symbols lie strictly on this plane, "
        "guaranteeing flicker-free illumination."
    )
    add_figure_with_caption("figures/fig1_csk_constellation_3d.png", "Figure 3.1: IEEE 802.15.7 4-CSK Constellation in Normalized 3D RGB Space (Sum I = 1.0)", width_in=5.8)

    add_heading_1("3.3 Channel Impairment Formulations")
    add_p(
        "The physical optical channel introduces three concurrent real-world sources of degradation modeled by the vector relation:"
    )
    make_equation_box(doc, "I_{rx} = alpha * ( M_{crosstalk} * I_{tx} ) + n_{AWGN}", "(3.1)")
    add_p("where I_tx is the (M x 3) transmitted symbol matrix and I_rx is the resulting received optical signal.")

    add_heading_2("3.3.1 Additive White Gaussian Noise (Thermal and Ambient Shot Noise)")
    add_p(
        "Optical receivers are corrupted by thermal noise originating from the photodiode load resistor and pre-amplifier electronics, "
        "alongside shot noise induced by ambient sunlight and fluorescent illumination. The noise vector n_{AWGN} is modeled as "
        "independent and identically distributed (i.i.d.) Gaussian random variables with zero mean and variance sigma_n^2:"
    )
    make_equation_box(doc, "sigma_n^2 = P_{nominal} / 10^{ (SNR_{dB} / 10) }", "(3.2)")
    add_p("where P_{nominal} is the average transmitted symbol power per channel: P_{nominal} = (1/3M) * sum || I_{tx} ||^2.")

    add_heading_2("3.3.2 LED Dimming Control and Photodiode Power Attenuation")
    add_p(
        "Indoor lighting installations permit users to adjust room illumination. In the simulator, the dimming factor alpha in [0.1, 1.0] "
        "linearly scales transmitted optical amplitude. As optical power is attenuated, the noise floor remains determined by ambient "
        "illumination, directly lowering effective channel SNR."
    )

    add_heading_2("3.3.3 Spectral Overlap and 3x3 RGB Optical Crosstalk Modeling")
    add_p(
        "Receiver color photodiodes use optical interference filters that possess overlapping spectral responsivity bands. For instance, "
        "a silicon photodiode filtered for the red band exhibits secondary sensitivity to green wavelengths. This phenomenon is modeled "
        "via a 3x3 optical crosstalk matrix M_{crosstalk}:"
    )
    make_equation_box(doc, "M_{crosstalk} = [ [ 0.90,  0.10,  0.00 ],  [ 0.10,  0.80,  0.10 ],  [ 0.00,  0.10,  0.90 ] ]", "(3.3)")
    add_p("which mixes adjacent optical channels and shifts received constellation centers away from ideal anchor coordinates.")

    add_heading_1("3.4 Maximum Likelihood Minimum Euclidean Distance Demapping")
    add_p(
        "At the receiver, 4-CSK demapping is performed using the Maximum Likelihood (ML) minimum Euclidean distance criterion across "
        "the reference constellation S_{ref} = { s_0, s_1, s_2, s_3 }:"
    )
    make_equation_box(doc, "hat{s}_i = argmin_{s in S_{ref}} || I_{rx,i} - s ||^2", "(3.4)")
    add_p("Each decoded symbol is subsequently mapped back into its corresponding 2-bit binary representation.")

    add_heading_1("3.5 Receiver Automatic Gain Control (AGC) Dynamics")
    add_p(
        "Because LED dimming scales the optical amplitude of incoming symbols, a standard demapper comparing received vectors against "
        "the fixed nominal constellation S_{ref} would misclassify symbols. To resolve this, the simulated receiver employs an Automatic "
        "Gain Control (AGC) stage that restores the nominal constellation scale:"
    )
    make_equation_box(doc, "I_{rx, AGC} = I_{rx} / alpha", "(3.5)")
    add_p(
        "Dividing by alpha restores the mean signal amplitude to nominal levels, but amplifies the noise variance by 1/alpha^2. This "
        "models the physical reality that dimmed optical links suffer elevated symbol error rates even when gain is electronically compensated."
    )

    print("Chapters 1 to 3 written successfully.")
